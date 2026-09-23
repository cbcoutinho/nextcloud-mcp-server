"""Unit tests for the .docx tracked-change writer (utils/docx_revisions.py).

Documents are built with python-docx and the result is inspected at the XML
level, since what matters is the exact ``<w:ins>``/``<w:del>`` markup editors
read, and re-opened with python-docx to prove the package is still valid.
"""

import io
import zipfile
from datetime import datetime, timezone

import docx
import pytest
from lxml import etree

from nextcloud_mcp_server.utils.docx_revisions import (
    W_NS,
    DocxRevisionError,
    apply_revision,
)

pytestmark = pytest.mark.unit

NS = {"w": W_NS}
FIXED_NOW = datetime(2026, 9, 23, 10, 0, 0, tzinfo=timezone.utc)


def _docx(*paragraphs: list[tuple[str, bool]] | str) -> bytes:
    """Build a .docx; a paragraph is a string, or a list of (text, bold) runs."""
    document = docx.Document()
    for spec in paragraphs:
        paragraph = document.add_paragraph()
        for text, bold in [(spec, False)] if isinstance(spec, str) else spec:
            paragraph.add_run(text).bold = bold
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _body(content: bytes) -> etree._Element:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return etree.fromstring(archive.read("word/document.xml"))


def _visible_text(paragraph: etree._Element) -> str:
    """Text a reader sees with all changes accepted."""
    return "".join(
        t.text or ""
        for t in paragraph.iter(f"{{{W_NS}}}t")
        if not any(a.tag == f"{{{W_NS}}}del" for a in t.iterancestors())
    )


def _apply(content: bytes, anchor: str, new_text: str = "", **kwargs):
    kwargs.setdefault("author", "alice")
    kwargs.setdefault("now", FIXED_NOW)
    return apply_revision(content, anchor, kwargs.pop("new_text", new_text), **kwargs)


def test_insert_wraps_new_text_in_w_ins_after_the_anchor():
    result = _apply(_docx("The quick fox."), "quick", " brown")

    body = _body(result.content)
    (ins,) = body.findall(".//w:ins", NS)
    assert ins.get(f"{{{W_NS}}}author") == "alice"
    assert ins.get(f"{{{W_NS}}}date") == "2026-09-23T10:00:00Z"
    assert int(ins.get(f"{{{W_NS}}}id")) == result.revision_ids[0]
    assert ins.findtext(".//w:t", namespaces=NS) == " brown"
    (paragraph,) = [p for p in body.iter(f"{{{W_NS}}}p") if _visible_text(p)]
    assert _visible_text(paragraph) == "The quick brown fox."
    assert result.paragraph_text == "The quick fox."
    docx.Document(io.BytesIO(result.content))  # still a valid package


def test_delete_moves_anchor_runs_into_w_del_as_deltext():
    result = _apply(_docx("Keep this, drop that."), " drop that", mode="delete")

    body = _body(result.content)
    (deletion,) = body.findall(".//w:del", NS)
    assert [t.text for t in deletion.iter(f"{{{W_NS}}}delText")] == [" drop that"]
    assert not deletion.findall(".//w:t", NS)
    (paragraph,) = [p for p in body.iter(f"{{{W_NS}}}p") if _visible_text(p)]
    assert _visible_text(paragraph) == "Keep this,."


def test_replace_emits_a_deletion_followed_by_an_insertion():
    result = _apply(_docx("Pay 10 EUR."), "10", "12", mode="replace")

    body = _body(result.content)
    deletion = body.find(".//w:del", NS)
    assert deletion is not None
    insertion = deletion.getnext()
    assert insertion.tag == f"{{{W_NS}}}ins"
    assert len(set(result.revision_ids)) == 2
    (paragraph,) = [p for p in body.iter(f"{{{W_NS}}}p") if _visible_text(p)]
    assert _visible_text(paragraph) == "Pay 12 EUR."


def test_anchor_spanning_differently_formatted_runs():
    """An anchor crossing run (formatting) boundaries is matched and split."""
    content = _docx([("Hello ", False), ("bold", True), (" world", False)])

    result = _apply(content, "lo bold wo", mode="delete")

    body = _body(result.content)
    deleted = [t.text for t in body.iter(f"{{{W_NS}}}delText")]
    assert "".join(deleted) == "lo bold wo"
    # The bold run keeps its formatting inside the deletion.
    assert body.find(".//w:del/w:r/w:rPr/w:b", NS) is not None
    (paragraph,) = [p for p in body.iter(f"{{{W_NS}}}p") if _visible_text(p)]
    assert _visible_text(paragraph) == "Helrld"


def test_insert_inherits_the_formatting_of_the_anchor_end():
    content = _docx([("plain ", False), ("bold", True)])

    result = _apply(content, "bold", "er")

    assert _body(result.content).find(".//w:ins/w:r/w:rPr/w:b", NS) is not None


def test_revision_id_does_not_collide_with_existing_ids():
    first = _apply(_docx("one two"), "one", "!")
    second = _apply(first.content, "two", "?")

    assert second.revision_ids[0] > first.revision_ids[0]


def test_other_package_parts_are_copied_unchanged():
    original = _docx("text")
    result = _apply(original, "text", "!")

    with (
        zipfile.ZipFile(io.BytesIO(original)) as before,
        zipfile.ZipFile(io.BytesIO(result.content)) as after,
    ):
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "word/document.xml":
                assert before.read(name) == after.read(name), name


def test_newline_and_tab_in_new_text_become_br_and_tab():
    result = _apply(_docx("a"), "a", "x\ny\tz")

    run = _body(result.content).find(".//w:ins/w:r", NS)
    tags = [etree.QName(c).localname for c in run if c.tag != f"{{{W_NS}}}rPr"]
    assert tags == ["t", "br", "t", "tab", "t"]


def test_surrounding_whitespace_is_preserved_in_split_runs():
    result = _apply(_docx("alpha beta"), "alpha", mode="delete")

    body = _body(result.content)
    remaining = [
        t for t in body.iter(f"{{{W_NS}}}t") if (t.text or "").strip() == "beta"
    ]
    assert remaining[0].text == " beta"
    assert remaining[0].get("{http://www.w3.org/XML/1998/namespace}space") == (
        "preserve"
    )


def test_proof_error_marks_between_runs_do_not_block_a_match():
    """Word splits words around w:proofErr; the anchor must still match."""
    original = _docx([("mis", False), ("spelt", True)])
    with zipfile.ZipFile(io.BytesIO(original)) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        parts = {n: archive.read(n) for n in archive.namelist()}
    first_run = root.find(".//w:r", NS)
    first_run.addnext(etree.Element(f"{{{W_NS}}}proofErr"))
    parts["word/document.xml"] = etree.tostring(root)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)

    result = _apply(out.getvalue(), "misspelt", mode="delete")

    deletion = _body(result.content).find(".//w:del", NS)
    assert "".join(t.text for t in deletion.iter(f"{{{W_NS}}}delText")) == "misspelt"
    # Document order is kept: the proofErr moved into the deletion with its runs.
    assert deletion.find("w:proofErr", NS) is not None


def test_ambiguous_anchor_is_refused_without_occurrence():
    content = _docx("same", "same")

    with pytest.raises(DocxRevisionError, match="occurs 2 times"):
        _apply(content, "same", "!")


def test_occurrence_selects_the_nth_match():
    content = _docx("same first", "same second")

    result = _apply(content, "same", "!", occurrence=2)

    assert result.paragraph_text == "same second"
    assert result.match_count == 2


def test_occurrence_out_of_range_is_refused():
    content = _docx("once")

    with pytest.raises(DocxRevisionError, match="occurs only 1"):
        _apply(content, "once", "!", occurrence=2)


def test_missing_anchor_is_refused():
    content = _docx("hello")

    with pytest.raises(DocxRevisionError, match="not found"):
        _apply(content, "absent", "!")


def test_anchor_does_not_cross_paragraphs():
    content = _docx("end of one", "start of two")

    with pytest.raises(DocxRevisionError, match="not found"):
        _apply(content, "one start", "!")


def test_anchor_does_not_cross_an_existing_tracked_change():
    first = _apply(_docx("left right"), "left", " middle")

    # "left middle right" is visible, but "middle" sits inside a w:ins.
    with pytest.raises(DocxRevisionError, match="not found"):
        _apply(first.content, "left middle", mode="delete")


@pytest.mark.parametrize(
    ("mode", "new_text", "message"),
    [
        ("delete", "x", "must be empty"),
        ("insert", "", "is required"),
        ("replace", "", "is required"),
    ],
)
def test_mode_and_new_text_must_agree(mode, new_text, message):
    content = _docx("text")

    with pytest.raises(DocxRevisionError, match=message):
        _apply(content, "text", new_text, mode=mode)


@pytest.mark.parametrize("field", ["new_text", "author"])
def test_xml_illegal_characters_are_refused(field):
    content = _docx("text")
    kwargs = {"new_text": "ok", "author": "alice", field: "bad\x07"}

    with pytest.raises(DocxRevisionError, match=f"{field} contains a control"):
        _apply(content, "text", **kwargs)


def test_empty_anchor_is_refused():
    content = _docx("text")

    with pytest.raises(DocxRevisionError, match="must not be empty"):
        _apply(content, "", "!")


def test_non_zip_input_is_refused():
    with pytest.raises(DocxRevisionError, match="not a zip"):
        _apply(b"plain text", "plain", "!")


def test_zip_that_is_not_a_word_document_is_refused():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("readme.txt", "hi")
    content = out.getvalue()

    with pytest.raises(DocxRevisionError, match="Not a Word document"):
        _apply(content, "hi", "!")
