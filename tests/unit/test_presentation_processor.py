"""Presentations are read shape-by-shape, keeping slide/table structure."""

import io

import pytest
from pptx import Presentation
from pptx.util import Inches

from nextcloud_mcp_server.document_processors.presentation import (
    PPTX_MIME,
    SLIDE_BOUNDARIES_KEY,
    PptxProcessor,
)

pytestmark = pytest.mark.unit


def _deck(slides: list[dict]) -> bytes:
    """A .pptx built from ``[{"title": str, "body": [str, ...], "table": [[...]]}]``."""
    prs = Presentation()
    for spec in slides:
        if "table" in spec:
            slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
            rows = spec["table"]
            shape = slide.shapes.add_table(
                len(rows), len(rows[0]), Inches(1), Inches(1), Inches(4), Inches(2)
            )
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    shape.table.cell(r, c).text = value
            continue

        slide = prs.slides.add_slide(prs.slide_layouts[1])
        if "title" in spec:
            slide.shapes.title.text = spec["title"]
        for line in spec.get("body", []):
            tf = slide.placeholders[1].text_frame
            para = tf.paragraphs[0] if not tf.text else tf.add_paragraph()
            para.text = line
        if "notes" in spec:
            slide.notes_slide.notes_text_frame.text = spec["notes"]

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_supported_mime_type_is_pptx_only():
    assert PptxProcessor().supported_mime_types == {PPTX_MIME}


async def test_text_frames_become_a_markdown_section():
    content = _deck([{"title": "Intro", "body": ["Line one", "Line two"]}])

    result = await PptxProcessor().process(content, PPTX_MIME, "d.pptx")

    assert "## Slide 1" in result.text
    assert "Intro" in result.text
    assert "Line one" in result.text
    assert "Line two" in result.text
    assert result.metadata["slide_count"] == 1
    assert result.metadata["parse_mode"] == "markdown"


async def test_a_table_shape_becomes_a_markdown_table():
    content = _deck([{"table": [["h1", "h2"], ["a", "b"]]}])

    result = await PptxProcessor().process(content, PPTX_MIME, "t.pptx")

    assert "| h1 | h2 |" in result.text
    assert "| a | b |" in result.text


async def test_pipe_in_a_cell_does_not_break_the_row():
    content = _deck([{"table": [["head"], ["a|b"]]}])

    result = await PptxProcessor().process(content, PPTX_MIME, "p.pptx")

    assert r"| a\|b |" in result.text


async def test_newline_in_a_cell_does_not_break_the_table():
    content = _deck([{"table": [["head"], ["line one\nline two"]]}])

    result = await PptxProcessor().process(content, PPTX_MIME, "n.pptx")

    assert "| line one line two |" in result.text


async def test_speaker_notes_are_included():
    content = _deck([{"title": "S", "body": ["x"], "notes": "Remember to smile"}])

    result = await PptxProcessor().process(content, PPTX_MIME, "notes.pptx")

    assert "**Notes:** Remember to smile" in result.text


async def test_each_slide_gets_its_own_boundary_span():
    content = _deck(
        [{"title": "First", "body": ["a"]}, {"title": "Second", "body": ["b"]}]
    )

    result = await PptxProcessor().process(content, PPTX_MIME, "two.pptx")

    spans = result.metadata[SLIDE_BOUNDARIES_KEY]
    assert [s["slide"] for s in spans] == [1, 2]
    # Offsets must index the returned text exactly -- they are what attributes a
    # chunk back to a slide, the presentation stand-in for a page number.
    for span in spans:
        segment = result.text[span["start_offset"] : span["end_offset"]]
        assert segment.startswith(f"## Slide {span['slide']}")
    assert spans[0]["end_offset"] == spans[1]["start_offset"]


async def test_empty_slide_is_skipped_not_emitted_as_an_empty_section():
    content = _deck([{}, {"title": "Real", "body": ["x"]}])

    result = await PptxProcessor().process(content, PPTX_MIME, "e.pptx")

    assert "## Slide 1" not in result.text
    assert "## Slide 2" in result.text
    assert [s["slide"] for s in result.metadata[SLIDE_BOUNDARIES_KEY]] == [2]


async def test_health_check_is_true_once_pptx_is_importable():
    assert await PptxProcessor().health_check() is True
